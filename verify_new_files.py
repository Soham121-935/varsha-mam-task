#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Quality-control for the two additional formatted workbooks.

Checks that every value survived, that no designation was invented, that all
highlighting is gone, that the master's design system was reproduced, and that
the two live formulas still work.
"""

import collections
import os
import re
import sys
import zipfile

import openpyxl
from PIL import ImageFont

from designations import build as build_designations
from report_format import GLYPH_FACTOR, CELL_PADDING, PX_PER_CHAR, squash

HERE = os.path.dirname(os.path.abspath(__file__))
DESIGN_SRC = os.path.join(HERE, "Summary AiFormative_and_Summative_Feedback_Report 2025-2026.xlsx")
MASTER25 = os.path.join(HERE, "Summary_AiFormative_and_Summative_Feedback_Report_2025-2026_Formatted.xlsx")

# sheet -> (src file, out file, rows, cols kept from src, name col in src, dropped cols)
PLAN = [
    dict(sheet="Formative", src="2023-24 (1).xlsx", out="2023-24 (1)_Formatted.xlsx",
         rows=(2, 49), cols=[1, 2, 3, 4, 6, 7], name_col=2, dropped=[5],
         title="AY_2023_24_Sem_II_Formative"),
    dict(sheet="Summative", src="2023-24 (1).xlsx", out="2023-24 (1)_Formatted.xlsx",
         rows=(1, 49), cols=[1, 2, 3, 4, 6, 7], name_col=2, dropped=[5],
         title="AY_2023_24_Sem_II_Summative"),
    dict(sheet="2024-25 Sem-I Formative", src="SY A 24-25 ODD (2).xlsx", out="SY A 24-25 ODD (2)_Formatted.xlsx",
         rows=(1, 71), cols=[1, 2, 3, 4, 6], name_col=2, dropped=[5],
         title="AY_2024_25_Sem_I_Formative"),
    dict(sheet="Sem-I Summative", src="SY A 24-25 ODD (2).xlsx", out="SY A 24-25 ODD (2)_Formatted.xlsx",
         rows=(1, 70), cols=[2, 3, 4, 6], name_col=2, dropped=[1, 5],
         title="AY_2024_25_Sem_I_Summative"),
    dict(sheet="Sem-II Formative", src="SY A 24-25 ODD (2).xlsx", out="SY A 24-25 ODD (2)_Formatted.xlsx",
         rows=(1, 133), cols=[1, 2, 3, 4, 5, 6], name_col=2, dropped=[],
         title="AY_2024_25_Sem_II_Formative"),
    dict(sheet="Sem-II Summative", src="SY A 24-25 ODD (2).xlsx", out="SY A 24-25 ODD (2)_Formatted.xlsx",
         rows=(5, 137), cols=[1, 2, 3, 4], name_col=1, dropped=[],
         title="AY_2024_25_Sem_II_Summative"),
]

failures, notes = [], []


def check(cond, msg):
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        failures.append(msg)


def note(msg):
    print("  note " + msg)
    notes.append(msg)


def main():
    lookup = build_designations(DESIGN_SRC)
    m = openpyxl.load_workbook(MASTER25)["summative"]

    print("\n[1] every value survived the rebuild")
    for p in PLAN:
        src_ws = openpyxl.load_workbook(os.path.join(HERE, p["src"]))[p["sheet"]]
        out_ws = openpyxl.load_workbook(os.path.join(HERE, p["out"]))[p["sheet"]]
        r0, r1 = p["rows"]
        original = [[src_ws.cell(row=r, column=c).value for c in p["cols"]]
                    for r in range(r0, r1 + 1)]
        i_name = p["cols"].index(p["name_col"])         # 0-based, before injection
        ncols = len(p["cols"]) + 1                      # + the Designation column
        got_rows = [[out_ws.cell(row=5 + i, column=1 + j).value for j in range(ncols)]
                    for i in range(len(original))]
        # strip the injected Designation column (0-based i_name + 1) back out
        got = [row[:i_name + 1] + row[i_name + 2:] for row in got_rows]
        check(len(original) == len(got), "%s: %d rows" % (p["sheet"], len(got)))
        bad = [(i, j, original[i][j], got[i][j])
               for i in range(min(len(original), len(got)))
               for j in range(len(p["cols"])) if original[i][j] != got[i][j]]
        check(not bad, "%s: all %d values identical to the original" % (p["sheet"], len(original) * len(p["cols"])))
        for b in bad[:5]:
            print("        ", b)

    print("\n[2] dropped columns really were constant / empty")
    for p in PLAN:
        if not p["dropped"]:
            continue
        src_ws = openpyxl.load_workbook(os.path.join(HERE, p["src"]))[p["sheet"]]
        r0, r1 = p["rows"]
        for c in p["dropped"]:
            vals = {str(src_ws.cell(row=r, column=c).value).strip()
                    for r in range(r0, r1 + 1) if src_ws.cell(row=r, column=c).value is not None}
            what = "constant %r" % vals.pop() if len(vals) == 1 else ("empty" if not vals else repr(vals))
            check(len(vals) <= 1, "%s: dropped column %s was %s" %
                  (p["sheet"], openpyxl.utils.get_column_letter(c), what))

    print("\n[3] designations: filled only from the source, nothing invented")
    for p in PLAN:
        out_ws = openpyxl.load_workbook(os.path.join(HERE, p["out"]))[p["sheet"]]
        i_name = p["cols"].index(p["name_col"])
        name_col_1 = i_name + 1                          # 1-based
        desig_col_1 = i_name + 2
        seen = collections.Counter()
        wrong = []
        for i in range(p["rows"][1] - p["rows"][0] + 1):
            row = 5 + i
            name = out_ws.cell(row=row, column=name_col_1).value
            got = out_ws.cell(row=row, column=desig_col_1).value
            want = lookup.get(squash(name)) if isinstance(name, str) else None
            seen[got if got is not None else "(blank)"] += 1
            if got != want:
                wrong.append((row, name, got, want))
        check(not wrong, "%s: designation column correct for all rows %s" % (p["sheet"], dict(seen)))
        for w in wrong[:5]:
            print("        ", w)

    print("\n[4] no highlighting left anywhere")
    for out in sorted({p["out"] for p in PLAN}):
        wb = openpyxl.load_workbook(os.path.join(HERE, out))
        filled = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.fill is not None and c.fill.fill_type not in (None, "none"):
                        filled.append((ws.title, c.coordinate, c.fill.fill_type,
                                       str(c.fill.fgColor.rgb)))
        check(not filled, "%s: every cell unfilled (%d filled cells)" % (out, len(filled)))
        for f in filled[:5]:
            print("        ", f)

    print("\n[5] master design system reproduced")
    ref = dict(
        title=(m["A1"].font.name, m["A1"].font.size, bool(m["A1"].font.bold)),
        header=(m["A4"].font.name, m["A4"].font.size, bool(m["A4"].font.bold)),
        data=(m["A5"].font.name, m["A5"].font.size, bool(m["A5"].font.bold)),
        side=m["A5"].border.left.style,
        h1=m.row_dimensions[1].height, h4=m.row_dimensions[4].height,
        zoom=m.sheet_view.zoomScale,
    )
    for p in PLAN:
        ws = openpyxl.load_workbook(os.path.join(HERE, p["out"]))[p["sheet"]]
        ncols = len(p["cols"]) + 1                       # table width (extras aside)
        nrows = p["rows"][1] - p["rows"][0] + 1
        check(ws["A1"].value == m["A1"].value and ws["A2"].value == m["A2"].value,
              "%s: college / department rows" % p["sheet"])
        check(ws["A3"].value == p["title"], "%s: title row = %s" % (p["sheet"], p["title"]))
        check(sorted(str(r) for r in ws.merged_cells.ranges) ==
              ["A1:%s1" % chr(64 + ncols), "A2:%s2" % chr(64 + ncols), "A3:%s3" % chr(64 + ncols)],
              "%s: title rows merged A1:%s3" % (p["sheet"], chr(64 + ncols)))
        check((ws["A1"].font.name, ws["A1"].font.size, bool(ws["A1"].font.bold)) == ref["title"],
              "%s: title font == master" % p["sheet"])
        hdr_ok = all(
            (ws.cell(4, c).font.name, ws.cell(4, c).font.size, bool(ws.cell(4, c).font.bold))
            == ref["header"]
            and all(getattr(ws.cell(4, c).border, s).style == ref["side"]
                    for s in ("left", "right", "top", "bottom"))
            and ws.cell(4, c).value not in (None, "")
            for c in range(1, ncols + 1))
        check(hdr_ok, "%s: header row bold + boxed, no empty labels" % p["sheet"])
        ok = True
        for r in range(5, 5 + nrows):
            for c in range(1, ncols + 1):
                cell = ws.cell(r, c)
                if not all(getattr(cell.border, s).style == ref["side"]
                           for s in ("left", "right", "top", "bottom")):
                    ok = False
                if (cell.font.name, cell.font.size, bool(cell.font.bold)) != ref["data"]:
                    ok = False
        check(ok, "%s: every one of %d data cells boxed, Calibri 12" % (p["sheet"], nrows * ncols))
        check(ws.row_dimensions[1].height == ref["h1"]
              and all(ws.row_dimensions[r].height == ref["h4"] for r in range(4, 5 + nrows)),
              "%s: row heights == master (%s / %s)" % (p["sheet"], ref["h1"], ref["h4"]))
        check(ws.sheet_view.zoomScale == ref["zoom"], "%s: zoom == master" % p["sheet"])
        check(len(ws._images) == 1, "%s: logo present" % p["sheet"])
        if ws._images:
            im = ws._images[0]
            check((im.anchor._from.col, im.anchor._from.row,
                   im.anchor._from.colOff, im.anchor._from.rowOff) == (0, 0, 130630, 64355),
                  "%s: logo anchored like the master" % p["sheet"])
            check((im.anchor.ext.width, im.anchor.ext.height) == (1415142, 577902),
                  "%s: logo same size (no distortion)" % p["sheet"])

    print("\n[6] the two live formulas survived and still point at Performance")
    ws = openpyxl.load_workbook(os.path.join(HERE, "SY A 24-25 ODD (2)_Formatted.xlsx"))["Sem-II Summative"]
    found = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.startswith("=")]
    check(sorted(f for _, f in found) == ["=ROUNDUP(E136,0)", "=ROUNDUP(E137,0)"],
          "formulas present and re-pointed to the new Performance column: %s" % sorted(found))
    check(ws.cell(136, 5).value is not None and ws.cell(137, 5).value is not None,
          "the cells the formulas read (E136/E137) hold the Performance values "
          "(%.2f / %.2f -> %s / %s)" % (ws.cell(136, 5).value, ws.cell(137, 5).value,
                                        round(ws.cell(136, 5).value + 0.4999),
                                        round(ws.cell(137, 5).value + 0.4999)))
    check(sorted(k for k, _ in found) == ["F136", "F137"], "formulas sit in F136/F137: %s"
          % sorted(k for k, _ in found))

    print("\n[7] nothing clipped, no dead whitespace")
    f12 = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 16)
    fb12 = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 16)
    for p in PLAN:
        ws = openpyxl.load_workbook(os.path.join(HERE, p["out"]))[p["sheet"]]
        nrows = p["rows"][1] - p["rows"][0] + 1
        over = []
        for c in range(1, len(p["cols"]) + 2):
            width = ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 9.0
            avail = width * PX_PER_CHAR + 5
            longest = 0
            for r in range(4, 5 + nrows):
                v = ws.cell(r, c).value
                if v is not None:
                    longest = max(longest, (fb12 if r == 4 else f12).getlength(str(v)))
            need = longest * GLYPH_FACTOR + CELL_PADDING
            if need > avail + 0.5 or avail - need > 45:
                over.append((openpyxl.utils.get_column_letter(c), round(need), round(avail)))
        check(not over, "%s: column widths fit %s" % (p["sheet"], over))

    print("\n[8] package hygiene")
    for out in sorted({p["out"] for p in PLAN}):
        path = os.path.join(HERE, out)
        with zipfile.ZipFile(path) as z:
            check(z.testzip() is None, "%s: zip intact" % out)
            bad = []
            for n in z.namelist():
                if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
                    if re.search(r'<col\b[^>]*\sstyle="\d+"', z.read(n).decode("utf-8")):
                        bad.append(n)
            check(not bad, "%s: no invalid <col style=...>" % out)
            names = re.findall(r'<sheet name="([^"]+)"', z.read("xl/workbook.xml").decode())
            print("       sheets: %s" % names)

    print("\n" + "=" * 72)
    if failures:
        print("FAILED %d check(s):" % len(failures))
        for f in failures:
            print("  -", f)
        sys.exit(1)
    print("ALL CHECKS PASSED")


if __name__ == "__main__":
    main()
