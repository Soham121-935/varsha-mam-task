#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Apply the master design system to the two additional workbooks
(``2023-24 (1).xlsx`` and ``SY A 24-25 ODD (2).xlsx``).

    python3 format_new_files.py

Design rules carried over from the 2025-26 deliverable:

* 3-row title block (college / department / report) merged across the table,
  then a bold header row, then the data - all boxed in the master's thin border
* the workbook's own logo, at the master's size and offset
* all highlighting removed - every cell is left unfilled
* a Designation column may be injected after Faculty Name, populated from the
  2025-26 workbook and left blank wherever that source is missing or disagrees.
  It is added per file - see WITH_DESIGNATION - and is off for both files: the
  2023-24 data is two academic years older than the source, so its designations
  would not be reliable.
* the sparse Sr. No. column may be dropped - see DROP_SR_NO. On request it is
  gone from the 24-25 file; the 2023-24 file keeps it.
"""

import os
import shutil

import openpyxl

from designations import build as build_designations
from report_format import (build_sheet, cleanup, extract_logo, fresh_sheet,
                           make_logo, read_design, reseat_images, shift_formula,
                           squash, tidy_package)

HERE = os.path.dirname(os.path.abspath(__file__))
DESIGN_SRC = os.path.join(HERE, "Summary AiFormative_and_Summative_Feedback_Report 2025-2026.xlsx")
BACKUP_DIR = os.path.join(HERE, "backup")

# Add the Designation column to this file? Keyed by source filename.
# Off for both: the 2023-24 reports are two academic years older than the only
# designation source (the 2025-26 workbook), so those designations would not be
# reliable, and the same column was then dropped from the 24-25 file on request.
WITH_DESIGNATION = {
    "2023-24 (1).xlsx":          False,
    "SY A 24-25 ODD (2).xlsx":   False,
}

# Drop the sparse Sr. No. column? Keyed by source filename.
# The Sr. No. column is only partly filled (41 of 71 rows on one sheet, 21 of
# 133 on the other) and restarts mid-table, so it carries no usable ordering.
DROP_SR_NO = {
    "2023-24 (1).xlsx":          False,
    "SY A 24-25 ODD (2).xlsx":   True,
}

FILES = [
    (os.path.join(HERE, "2023-24 (1).xlsx"),
     os.path.join(HERE, "2023-24 (1)_Formatted.xlsx")),
    (os.path.join(HERE, "SY A 24-25 ODD (2).xlsx"),
     os.path.join(HERE, "SY A 24-25 ODD (2)_Formatted.xlsx")),
]


# --------------------------------------------------------------------------- #
def grid(ws, rows, cols):
    r0, r1 = rows
    return [[ws.cell(row=r, column=c).value for c in cols] for r in range(r0, r1 + 1)]


def lone_value(ws, col, rows):
    """The single distinct non-empty value in a column, else None."""
    r0, r1 = rows
    vals = {str(ws.cell(row=r, column=col).value).strip()
            for r in range(r0, r1 + 1) if ws.cell(row=r, column=col).value is not None}
    return vals.pop() if len(vals) == 1 else None


def is_empty(ws, col, rows):
    r0, r1 = rows
    return all(ws.cell(row=r, column=col).value is None for r in range(r0, r1 + 1))


def with_designation(rows, name_idx, lookup):
    """Insert a Designation column straight after the Faculty Name column."""
    out = []
    for row in rows:
        name = row[name_idx] if name_idx < len(row) else None
        desig = lookup.get(squash(name)) if isinstance(name, str) else None
        out.append(row[:name_idx + 1] + [desig] + row[name_idx + 1:])
    return out


def headers_with(base, name_idx, with_desig):
    """Header row, with a Designation column added after Faculty Name."""
    if not with_desig:
        return list(base)
    return base[:name_idx + 1] + ["Designation"] + base[name_idx + 1:]


def numeric_idx(headers, names):
    """0-based indices of the given columns, looked up by header name."""
    return {headers.index(n) for n in names if n in headers}


# --------------------------------------------------------------------------- #
def build_file(src, out, design, lookup):
    key = os.path.basename(src)
    desig = WITH_DESIGNATION[key]
    srno = DROP_SR_NO[key]
    logo_png = extract_logo(src) or extract_logo(DESIGN_SRC)
    wb = openpyxl.load_workbook(src)
    reseat_images(wb, logo_png)

    # ------------------------------------------------ 2023-24 (1).xlsx ----
    if os.path.basename(src) == "2023-24 (1).xlsx":
        # ---- Formative: header on row 1, data rows 2-49, col E a constant
        old = wb["Formative"]
        rows = (2, 49)
        title = lone_value(old, 5, rows)
        assert title == "AY_2023_24_Sem_II_Formative", title
        data = grid(old, rows, [1, 2, 3, 4, 6, 7])          # A B C D F G  (E dropped)
        if desig:
            data = with_designation(data, 1, lookup)
        headers = headers_with(["Sr. No.", "Faculty Name", "Semester/ Group",
                                "Subject", "Attendees", "Feedback Percentage"], 1, desig)
        ws = fresh_sheet(wb, old)
        build_sheet(ws, design, title, headers, data,
                    numeric_cols=numeric_idx(headers, {"Sr. No.", "Attendees",
                                                       "Feedback Percentage"}),
                    logo=make_logo(logo_png))

        # ---- Summative: no header row, data rows 1-49, col E a constant
        old = wb["Summative"]
        rows = (1, 49)
        title = lone_value(old, 5, rows)
        assert title == "AY_2023_24_Sem_II_Summative", title
        data = grid(old, rows, [1, 2, 3, 4, 6, 7])
        if desig:
            data = with_designation(data, 1, lookup)
        ws = fresh_sheet(wb, old)
        # labels taken from the sibling Formative sheet, which has the same layout
        build_sheet(ws, design, title, headers, data,
                    numeric_cols=numeric_idx(headers, {"Sr. No.", "Attendees",
                                                       "Feedback Percentage"}),
                    logo=make_logo(logo_png))

    # ------------------------------------- SY A 24-25 ODD (2).xlsx --------
    else:
        # ---- 2024-25 Sem-I Formative: data rows 1-71, col E a constant
        old = wb["2024-25 Sem-I Formative"]
        rows = (1, 71)
        title = lone_value(old, 5, rows)
        assert title == "AY_2024_25_Sem_I_Formative", title
        keep = [2, 3, 4, 6] if srno else [1, 2, 3, 4, 6]    # A dropped with Sr. No.
        data = grid(old, rows, keep)                        # E dropped either way
        if desig:
            data = with_designation(data, 0 if srno else 1, lookup)
        headers = headers_with(["Faculty Name", "Sem/Class", "Course", "Performance"]
                               if srno else
                               ["Sr. No.", "Faculty Name", "Sem/Class",
                                "Course", "Performance"], 0 if srno else 1, desig)
        ws = fresh_sheet(wb, old)
        build_sheet(ws, design, title, headers, data,
                    numeric_cols=numeric_idx(headers, {"Sr. No.", "Performance"}),
                    logo=make_logo(logo_png))

        # ---- Sem-I Summative: data rows 1-70, col A empty, col E a constant
        old = wb["Sem-I Summative"]
        rows = (1, 70)
        assert is_empty(old, 1, rows), "col A was expected to be empty"
        title = lone_value(old, 5, rows)
        assert title == "AY_2024_25_Sem_I_Summative", title
        data = grid(old, rows, [2, 3, 4, 6])                # B C D F  (A and E dropped)
        if desig:
            data = with_designation(data, 0, lookup)
        headers = headers_with(["Faculty Name", "Sem/Class", "Course",
                                "Performance"], 0, desig)
        ws = fresh_sheet(wb, old)
        build_sheet(ws, design, title, headers, data,
                    numeric_cols=numeric_idx(headers, {"Performance"}),
                    logo=make_logo(logo_png))

        # ---- Sem-II Formative: data rows 1-133, col E holds two values -> kept
        old = wb["Sem-II Formative"]
        rows = (1, 133)
        assert lone_value(old, 5, rows) is None, "col E was expected to vary"
        keep = [2, 3, 4, 5, 6] if srno else [1, 2, 3, 4, 5, 6]   # A = Sr. No.
        data = grid(old, rows, keep)
        if desig:
            data = with_designation(data, 0 if srno else 1, lookup)
        headers = headers_with(["Faculty Name", "Sem/Class", "Course",
                                "Feedback", "Performance"] if srno else
                               ["Sr. No.", "Faculty Name", "Sem/Class", "Course",
                                "Feedback", "Performance"], 0 if srno else 1, desig)
        ws = fresh_sheet(wb, old)
        # title follows the AY_..._Sem_X_Type convention of its sibling sheets
        build_sheet(ws, design, "AY_2024_25_Sem_II_Formative", headers, data,
                    numeric_cols=numeric_idx(headers, {"Sr. No.", "Performance"}),
                    logo=make_logo(logo_png))

        # ---- Sem-II Summative: already laid out; data rows 5-137, cols A-D,
        #      plus two live =ROUNDUP() formulas in column E
        old = wb["Sem-II Summative"]
        rows = (5, 137)
        title = old["A3"].value
        assert title == "AY_2024_25_Sem_II_Summative", title
        formulas = [(r, old.cell(row=r, column=5).value)
                    for r in range(rows[0], rows[1] + 1)
                    if isinstance(old.cell(row=r, column=5).value, str)
                    and old.cell(row=r, column=5).value.startswith("=")]
        assert len(formulas) == 2, formulas
        data = grid(old, rows, [1, 2, 3, 4])
        if desig:
            data = with_designation(data, 0, lookup)
        headers = headers_with(["Faculty Name", "Sem/Class", "Course",
                                "Performance"], 0, desig)
        # A=Faculty Name, B=Sem/Class, C=Course, D=Performance, so Performance
        # is at 0-based index 3 and shifts one right when Designation is added.
        # The formulas reference it and sit in the column right after it.
        perf = 3 + (1 if desig else 0)
        shift = 1 if desig else 0
        extra = [(r - rows[0], perf + 1, shift_formula(f, 4, shift), True)
                 for r, f in formulas]
        ws = fresh_sheet(wb, old)
        build_sheet(ws, design, title, headers, data,
                    numeric_cols=numeric_idx(headers, {"Performance"}),
                    logo=make_logo(logo_png), extra_cells=extra)

    wb.active = 0
    wb.save(out)
    stripped = tidy_package(out, None, preserve_dimension_of=None)
    return stripped


def main():
    os.makedirs(BACKUP_DIR, exist_ok=True)
    lookup = build_designations(DESIGN_SRC)

    master_wb = openpyxl.load_workbook(DESIGN_SRC)
    design = read_design(master_wb)

    for src, out in FILES:
        backup = os.path.join(BACKUP_DIR, os.path.basename(src))
        if not os.path.exists(backup):
            shutil.copy2(src, backup)
        stripped = build_file(src, out, design, lookup)
        print("wrote %s  (%d stray <col style=...> removed)" % (os.path.basename(out), stripped))
    cleanup()


if __name__ == "__main__":
    main()
